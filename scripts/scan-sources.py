#!/usr/bin/env python3
"""6 kaynak dosya + 2 XLSX'i bir kez tara, kalıcı şema kaydı üret."""
import json
import openpyxl
from pathlib import Path
from collections import Counter
from lxml import etree

SRC = Path.home() / "pipeline" / "input" / "gs-fb-2026-10-26"
OUT = Path.home() / "hpanalytic" / "docs" / "VERI-SOZLESMESI.json"
OUT_MD = Path.home() / "hpanalytic" / "docs" / "VERI-SOZLESMESI.md"

report = {"kaynak_klasor": str(SRC), "dosyalar": [], "ozet": {}}

# ── XML'LER ──
for xml_file in sorted(SRC.glob("*.xml")):
    try:
        tree = etree.parse(str(xml_file))
        root = tree.getroot()

        # Tüm element isimleri ve attribute'ları
        elem_names = Counter()
        attr_names = Counter()
        samples = []
        for el in root.iter():
            tag = etree.QName(el).localname
            if tag == 'document' or tag == 'XML':
                continue
            elem_names[tag] += 1
            for a in el.attrib:
                attr_names[a] += 1
            if len(samples) < 5 and len(el.attrib) > 1:
                samples.append({"tag": tag, "attrs": dict(el.attrib)})

        report["dosyalar"].append({
            "dosya": xml_file.name,
            "tip": "XML",
            "boyut_kb": round(xml_file.stat().st_size / 1024, 1),
            "root": etree.QName(root).localname,
            "element_sayilari": dict(elem_names.most_common(10)),
            "attribute_sayilari": dict(attr_names.most_common(20)),
            "ornek_5": samples,
        })
    except Exception as e:
        report["dosyalar"].append({"dosya": xml_file.name, "hata": str(e)})

# ── CSV'LER ──
import csv
for csv_file in sorted(SRC.glob("*.csv")):
    try:
        with open(csv_file, encoding='utf-8', errors='ignore') as f:
            # ilk satır header
            header_line = f.readline().strip()
            delimiter = ';' if ';' in header_line else ','
            f.seek(0)
            reader = csv.reader(f, delimiter=delimiter)
            rows = list(reader)
            header = rows[0] if rows else []
            data_rows = rows[1:]

            # Sütun istatistikleri
            col_stats = {}
            for i, col in enumerate(header[:50]):
                vals = [r[i] for r in data_rows if i < len(r) and r[i].strip()]
                unique = len(set(vals))
                sample = list(set(vals))[:3]
                col_stats[col] = {
                    "dolu": len(vals),
                    "tekil": unique,
                    "ornek": sample,
                }

            report["dosyalar"].append({
                "dosya": csv_file.name,
                "tip": "CSV",
                "delimiter": delimiter,
                "boyut_kb": round(csv_file.stat().st_size / 1024, 1),
                "satir": len(data_rows),
                "sutun_sayisi": len(header),
                "sutunlar": header[:60],
                "kolon_istatistikleri": col_stats,
            })
    except Exception as e:
        report["dosyalar"].append({"dosya": csv_file.name, "hata": str(e)})

# ── XLSX'LER ──
for xlsx_file in sorted(SRC.glob("*.xlsx")):
    try:
        wb = openpyxl.load_workbook(xlsx_file, data_only=True)
        sheets = []
        for sn in wb.sheetnames:
            ws = wb[sn]
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            headers = [h for h in headers if h]
            sample_rows = []
            for r in range(2, min(5, ws.max_row + 1)):
                row = {}
                for c in range(1, min(15, ws.max_column + 1)):
                    h = ws.cell(1, c).value
                    v = ws.cell(r, c).value
                    if h and v is not None:
                        row[h] = str(v)[:50]
                if row:
                    sample_rows.append(row)
            sheets.append({
                "sheet": sn,
                "satir": ws.max_row - 1,
                "sutun_sayisi": len(headers),
                "ilk_15_sutun": headers[:15],
                "ornek_ilk_satir": sample_rows[0] if sample_rows else {},
            })

        report["dosyalar"].append({
            "dosya": xlsx_file.name,
            "tip": "XLSX",
            "boyut_kb": round(xlsx_file.stat().st_size / 1024, 1),
            "sheet_sayisi": len(sheets),
            "sheetler": sheets,
        })
    except Exception as e:
        report["dosyalar"].append({"dosya": xlsx_file.name, "hata": str(e)})

# ── ÖZET ──
report["ozet"] = {
    "toplam_dosya": len(report["dosyalar"]),
    "xml": sum(1 for d in report["dosyalar"] if d.get("tip") == "XML"),
    "csv": sum(1 for d in report["dosyalar"] if d.get("tip") == "CSV"),
    "xlsx": sum(1 for d in report["dosyalar"] if d.get("tip") == "XLSX"),
}

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2))

# Markdown özet
md = ["# Veri Sözleşmesi — GS-FB Maçı Kaynak Dosyaları\n",
      f"**Klasör:** `{SRC}`\n",
      f"**Toplam:** {report['ozet']['toplam_dosya']} dosya "
      f"({report['ozet']['xml']} XML · {report['ozet']['csv']} CSV · {report['ozet']['xlsx']} XLSX)\n"]
for d in report["dosyalar"]:
    md.append(f"\n## {d['dosya']}")
    md.append(f"- **Tip:** {d.get('tip', '?')} · **Boyut:** {d.get('boyut_kb', '?')} KB")
    if d.get("tip") == "XML":
        md.append(f"- **Root:** `{d.get('root')}`")
        md.append(f"- **Element dağılımı:** {d.get('element_sayilari', {})}")
        md.append(f"- **Attribute'lar:** {list(d.get('attribute_sayilari', {}).keys())[:15]}")
    elif d.get("tip") == "CSV":
        md.append(f"- **Satır:** {d.get('satir')} · **Sütun:** {d.get('sutun_sayisi')} · **Delimiter:** `{d.get('delimiter')}`")
        md.append(f"- **Sütunlar:** `{', '.join(d.get('sutunlar', [])[:20])}`")
    elif d.get("tip") == "XLSX":
        md.append(f"- **Sheet sayısı:** {d.get('sheet_sayisi')}")
        for sh in d.get("sheetler", []):
            md.append(f"  - `{sh['sheet']}` → {sh['satir']} satır × {sh['sutun_sayisi']} sütun")
            md.append(f"    - İlk 15: `{', '.join(str(x) for x in sh['ilk_15_sutun'][:15])}`")

OUT_MD.write_text("\n".join(md))
print(f"✓ JSON: {OUT}")
print(f"✓ MD: {OUT_MD}")
print(f"  {report['ozet']}")
