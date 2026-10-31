import openpyxl
from pathlib import Path

def build_xg_flow(match, xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    H = {ws.cell(1, c).value: c for c in range(1, ws.max_column+1) if ws.cell(1,c).value}

    player_xg = {}
    for row in range(2, ws.max_row+1):
        name = ws.cell(row, H['Player']).value
        if not name: continue
        xg = ws.cell(row, H['xG (expected goals)']).value
        try: player_xg[name] = float(xg) if xg not in (None, '-', '') else 0
        except: player_xg[name] = 0

    shot_count = {}
    for e in match.events:
        if e.kind in ('shot', 'goal') and e.player_name:
            shot_count[e.player_name] = shot_count.get(e.player_name, 0) + 1

    gs_flow, lil_flow = [], []
    gs_cum, lil_cum = 0, 0
    for e in sorted(match.events, key=lambda x: x.t):
        if e.kind not in ('shot', 'goal'): continue
        if e.player_name and shot_count.get(e.player_name, 0) > 0:
            per_shot = player_xg.get(e.player_name, 0) / shot_count[e.player_name]
        else:
            per_shot = 0.10
        entry = {'minute': round(e.t/60, 1), 'xg': round(per_shot, 3),
                 'is_goal': e.kind == 'goal', 'player': e.player_name}
        team = e.team_code()
        if team == 'GS':
            gs_cum += per_shot
            entry['cumulative'] = round(gs_cum, 2)
            gs_flow.append(entry)
        elif team == 'LIL':
            lil_cum += per_shot
            entry['cumulative'] = round(lil_cum, 2)
            lil_flow.append(entry)

    return {'GS': gs_flow, 'LIL': lil_flow,
            'totals': {'GS': round(gs_cum, 2), 'LIL': round(lil_cum, 2)},
            'method': 'XLSX oyuncu xG / oyuncu şut sayısı'}
